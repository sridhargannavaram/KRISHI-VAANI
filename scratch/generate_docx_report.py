import re
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_word_report():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Title
    title = doc.add_paragraph()
    title_run = title.add_run("KRISHI VAANI (कृषि वाणी)")
    title_run.font.size = Pt(24)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(21, 128, 61) # Primary Green
    title.paragraph_format.space_after = Pt(2)

    # Subtitle
    sub = doc.add_paragraph()
    sub_run = sub.add_run("Complete System Architecture, End-to-End User Flow, PostGIS Geospatial Engine, Multilingual AI Pipeline & Technical Verification Report")
    sub_run.font.size = Pt(11)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(71, 85, 105)
    sub.paragraph_format.space_after = Pt(14)

    # Metadata Callout Box / Table
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Target Platform:", "Production / Staging Readiness", "Database & GIS:", "Supabase PostgreSQL 15+ & PostGIS"),
        ("Audit Date:", "August 27, 2026", "Final Decision:", "READY FOR STAGING (100% Tests Passed)")
    ]
    
    for row_idx, data in enumerate(meta_data):
        row = meta_table.rows[row_idx]
        cell_0 = row.cells[0]
        cell_0.text = f"{data[0]} {data[1]}"
        cell_1 = row.cells[1]
        cell_1.text = f"{data[2]} {data[3]}"

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Read Markdown source
    md_path = "KRISHI_VAANI_COMPLETE_TECHNICAL_REPORT.md"
    if not os.path.exists(md_path):
        print(f"Error: {md_path} not found")
        return

    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    in_table = False
    table_rows = []
    in_code = False
    code_lines = []

    for line in lines:
        raw_line = line.rstrip("\n")
        stripped = raw_line.strip()

        # Handle Code Block
        if stripped.startswith("```"):
            if in_code:
                # End of code block
                in_code = False
                code_p = doc.add_paragraph()
                code_run = code_p.add_run("\n".join(code_lines))
                code_run.font.name = "Consolas"
                code_run.font.size = Pt(8.5)
                code_p.paragraph_format.space_after = Pt(8)
                code_lines = []
            else:
                in_code = True
                code_lines = []
            continue

        if in_code:
            code_lines.append(raw_line)
            continue

        # Handle Markdown Tables
        if "|" in stripped and stripped.startswith("|") and stripped.endswith("|"):
            # Check if separator row
            if re.match(r"^\|(\s*[-:]+\s*\|)+$", stripped):
                continue
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            table_rows.append(cells)
            in_table = True
            continue
        elif in_table:
            # End of table, write it to docx
            if table_rows:
                cols_count = max(len(r) for r in table_rows)
                t = doc.add_table(rows=len(table_rows), cols=cols_count)
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                for r_idx, r_data in enumerate(table_rows):
                    for c_idx, val in enumerate(r_data):
                        if c_idx < cols_count:
                            cell = t.cell(r_idx, c_idx)
                            # Strip markdown bolding/code marks
                            clean_val = val.replace("**", "").replace("`", "")
                            cell.text = clean_val
                            if r_idx == 0:
                                for p in cell.paragraphs:
                                    for run in p.runs:
                                        run.font.bold = True
                doc.add_paragraph().paragraph_format.space_after = Pt(6)
            table_rows = []
            in_table = False

        if not stripped:
            continue

        # Heading 1
        if stripped.startswith("# "):
            h_text = stripped[2:].replace("**", "").strip()
            h = doc.add_heading(h_text, level=1)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(4)
            continue

        # Heading 2
        if stripped.startswith("## "):
            h_text = stripped[3:].replace("**", "").strip()
            h = doc.add_heading(h_text, level=2)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(3)
            continue

        # Heading 3
        if stripped.startswith("### "):
            h_text = stripped[4:].replace("**", "").strip()
            h = doc.add_heading(h_text, level=3)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
            continue

        # Bullet lists
        if stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            b_text = stripped[2:].strip()
            parts = re.split(r'(\*\*.*?\*\*)', b_text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    p.add_run(part)
            p.paragraph_format.space_after = Pt(2)
            continue

        # Numbered lists
        num_match = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if num_match:
            p = doc.add_paragraph(style='List Number')
            n_text = num_match.group(2).strip()
            parts = re.split(r'(\*\*.*?\*\*)', n_text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    p.add_run(part)
            p.paragraph_format.space_after = Pt(2)
            continue

        # Standard Paragraph
        p = doc.add_paragraph()
        parts = re.split(r'(\*\*.*?\*\*)', stripped)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                r = p.add_run(part[2:-2])
                r.font.bold = True
            else:
                p.add_run(part)
        p.paragraph_format.space_after = Pt(4)

    output_path = "KRISHI_VAANI_COMPLETE_TECHNICAL_REPORT.docx"
    doc.save(output_path)
    print(f"Successfully compiled {output_path}")

if __name__ == "__main__":
    create_word_report()
